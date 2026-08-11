"""Tests for Status Comparison Report (Original Readiness vs Assigned Status).

Verifies the bug fix where the "Readiness Status" column previously showed the
AI-adjusted `detailed_status`. Now it must show the raw Readisys `device.status`.
"""
import io
import os
import pytest
import requests
from openpyxl import load_workbook

BASE_URL = (os.environ.get("REACT_APP_BACKEND_URL") or "https://cardiac-command.preview.emergentagent.com").rstrip("/")
API = f"{BASE_URL}/api"

ADMIN_USER = "futureroads"
ADMIN_PASS = "@@U1s9m6c7@@"


@pytest.fixture(scope="session")
def token():
    r = requests.post(
        f"{API}/auth/login",
        json={"username": ADMIN_USER, "password": ADMIN_PASS},
        timeout=90,
    )
    assert r.status_code == 200, f"login failed: {r.status_code} {r.text[:300]}"
    data = r.json()
    tok = data.get("access_token") or data.get("token")
    assert tok, f"no token in login response: {data}"
    return tok


@pytest.fixture(scope="session")
def auth_headers(token):
    return {"Authorization": f"Bearer {token}"}


# ── Subscribers list ────────────────────────────────────────────────────────
def test_subscribers_list_has_at_least_20(auth_headers):
    r = requests.get(f"{API}/reports/status-comparison/subscribers", headers=auth_headers, timeout=60)
    assert r.status_code == 200, r.text[:300]
    data = r.json()
    assert "subscribers" in data
    subs = data["subscribers"]
    assert isinstance(subs, list)
    assert len(subs) >= 20, f"expected >=20 subscribers, got {len(subs)}"


# ── Lawson State: sentinel_id ODWCB-S15 should have Original=UNKNOWN vs Assigned=REPOSITION
def test_lawson_state_odwcb_s15_original_unknown(auth_headers):
    sub = "Lawson State Community College"
    r = requests.get(
        f"{API}/reports/status-comparison",
        params={"subscriber": sub},
        headers=auth_headers,
        timeout=180,
    )
    assert r.status_code == 200, r.text[:300]
    data = r.json()
    assert data["subscriber"] == sub
    rows = data["rows"]
    assert isinstance(rows, list) and len(rows) > 0, "no rows"

    target = next((row for row in rows if row.get("sentinel_id") == "ODWCB-S15"), None)
    assert target is not None, f"ODWCB-S15 not found; sids={[r.get('sentinel_id') for r in rows][:10]}"

    print(f"ODWCB-S15 row: {target}")
    assert "readiness_status" in target
    assert "assigned_status" in target
    assert target["readiness_status"] == "UNKNOWN", (
        f"readiness_status should be UNKNOWN (original), got {target['readiness_status']!r}"
    )
    assert target["assigned_status"] == "REPOSITION", (
        f"assigned_status should be REPOSITION (AI-adjusted), got {target['assigned_status']!r}"
    )
    assert target["match"] is False, "match should be False when UNKNOWN != REPOSITION"


# ── Iberville Parish Sheriff: separate fields populated
def test_iberville_rows_have_separate_fields(auth_headers):
    sub = "Iberville Parish Sheriff"
    r = requests.get(
        f"{API}/reports/status-comparison",
        params={"subscriber": sub},
        headers=auth_headers,
        timeout=180,
    )
    assert r.status_code == 200, r.text[:300]
    data = r.json()
    rows = data["rows"]
    assert len(rows) > 0, "no rows returned for Iberville"

    # Every row must have both fields
    for row in rows:
        assert "readiness_status" in row
        assert "assigned_status" in row

    non_empty_readiness = [r for r in rows if r.get("readiness_status")]
    assert len(non_empty_readiness) >= 1, (
        f"expected >=1 row with non-empty readiness_status, got 0. "
        f"Sample rows: {rows[:3]}"
    )

    # summary sanity
    summary = data["summary"]
    assert summary["total"] == len(rows)
    computed_matches = sum(1 for row in rows if row["readiness_status"] == row["assigned_status"] and row["readiness_status"])
    # summary matches counts where match=True (which requires both set + equal)
    reported_matches = sum(1 for row in rows if row.get("match"))
    assert summary["matches"] == reported_matches
    computed_mismatches = sum(
        1 for row in rows
        if row["readiness_status"] and row["assigned_status"] and row["readiness_status"] != row["assigned_status"]
    )
    assert summary["mismatches"] == computed_mismatches, (
        f"summary.mismatches={summary['mismatches']} but computed={computed_mismatches}"
    )


# ── XLSX headers
def test_iberville_xlsx_headers(auth_headers):
    sub = "Iberville Parish Sheriff"
    r = requests.get(
        f"{API}/reports/status-comparison.xlsx",
        params={"subscriber": sub},
        headers=auth_headers,
        timeout=240,
    )
    assert r.status_code == 200, r.text[:300]
    ctype = r.headers.get("content-type", "")
    assert "spreadsheet" in ctype or "excel" in ctype or "octet-stream" in ctype, ctype
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    header_row = [ws.cell(row=1, column=i).value for i in range(1, 9)]
    expected = ["Model", "AED Serial", "Sentinel ID", "Original Readiness", "Assigned Status", "Match?", "Location", "Image"]
    assert header_row == expected, f"headers mismatch: {header_row}"
